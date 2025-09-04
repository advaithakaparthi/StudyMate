import os
from langchain.text_splitter import CharacterTextSplitter
from langchain_community.vectorstores import FAISS
from langchain_ollama import OllamaEmbeddings
from docx import Document
import fitz
import openpyxl

def extract_text(filepath):
    ext = os.path.splitext(filepath)[1].lower()
    if ext == '.pdf':
        return extract_text_from_pdf(filepath)
    elif ext == '.docx':
        return extract_text_from_docx(filepath)
    elif ext == '.xlsx':
        return extract_text_from_xlsx(filepath)
    elif ext == '.txt':
        return extract_text_from_txt(filepath)
    else:
        return ""

def extract_text_from_pdf(filepath):
    doc = fitz.open(filepath)
    text = ""
    for page in doc:
        text += page.get_text()
    return text

def extract_text_from_docx(filepath):
    doc = Document(filepath)
    return "\n".join([para.text for para in doc.paragraphs])

def extract_text_from_xlsx(filepath):
    wb = openpyxl.load_workbook(filepath, data_only=True)
    text = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            row_text = " ".join([str(cell) for cell in row if cell is not None])
            text.append(row_text)
    return "\n".join(text)

def extract_text_from_txt(filepath):
    with open(filepath, 'r', encoding='utf-8') as f:
        return f.read()

def process_documents(documents_folder, vector_store_path):
    texts = []
    for filename in os.listdir(documents_folder):
        file_path = os.path.join(documents_folder, filename)
        text = extract_text(file_path)
        if text:
            texts.append(text)
    full_text = "\n".join(texts)

    splitter = CharacterTextSplitter(chunk_size=500, chunk_overlap=50)
    chunks = splitter.split_text(full_text)

    embeddings = OllamaEmbeddings(model="mistral")
    vectorstore = FAISS.from_texts(chunks, embeddings)

    vectorstore.save_local(vector_store_path)
